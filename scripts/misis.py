#!/usr/bin/env python3
"""Read public MISIS timetables. No university credentials required.
Supports the observed Date/Number/Time and Day/Time exports; other layouts stay errors.
"""

import argparse
import concurrent.futures
import datetime
import hashlib
import io
import json
import re
import sys
import time
import urllib.request
from html.parser import HTMLParser
from pathlib import Path

PAGE = "https://misis.ru/students/schedule/"
DAYS = {
    "понедельник": 1,
    "вторник": 2,
    "среда": 3,
    "четверг": 4,
    "пятница": 5,
    "суббота": 6,
    "воскресенье": 0,
}


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self.current = None

    def handle_starttag(self, tag, attrs):
        u = dict(attrs).get("href", "")
        if tag == "a" and re.search(r"\.xlsx?$", u, re.I) and u.startswith("/files/"):
            self.current = {"url": "https://misis.ru" + u, "name": ""}

    def handle_data(self, data):
        if self.current is not None:
            self.current["name"] += data

    def handle_endtag(self, tag):
        if tag == "a" and self.current is not None:
            self.current["name"] = " ".join(self.current["name"].split())
            self.links.append(self.current)
            self.current = None


def download(url):
    """Fetch only official sources with bounded size, retries and no redirects."""
    if not url.startswith("https://misis.ru/"):
        raise ValueError("Разрешён только официальный сайт МИСИС")

    class NoRedirect(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, *args):
            raise ValueError("Неожиданное перенаправление источника")

    opener = urllib.request.build_opener(NoRedirect)
    last = None
    for attempt in range(3):
        try:
            with opener.open(
                urllib.request.Request(
                    url, headers={"User-Agent": "MISIS-Starosta/0.1", "Referer": PAGE}
                ),
                timeout=20,
            ) as r:
                data = r.read(12 * 1024 * 1024 + 1)
                if len(data) > 12 * 1024 * 1024:
                    raise ValueError("Слишком большой файл расписания")
                return data
        except Exception as exc:
            last = exc
            time.sleep(attempt)
    raise last


def sheets(data):
    if data[:2] == b"PK":
        import openpyxl

        w = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        for s in w:
            if s.max_row > 3000 or s.max_column > 2000:
                continue
            rows = [
                [str(v).strip() if v is not None else "" for v in row]
                for row in s.iter_rows(values_only=True)
            ]
            merges = [
                (r.min_row - 1, r.max_row, r.min_col - 1, r.max_col) for r in s.merged_cells.ranges
            ]
            yield s.title, rows, merges
    else:
        import xlrd

        w = xlrd.open_workbook(file_contents=data, formatting_info=True)
        for s in w.sheets():
            rows = [
                [str(v).strip() if v != "" else "" for v in s.row_values(r)] for r in range(s.nrows)
            ]
            yield s.name, rows, s.merged_cells


def parse_workbook(data):
    """Parse known spreadsheet layouts, preserving rejected cells for manual review."""
    groups = []
    warnings = []
    for sheet, rows, merges in sheets(data):
        # Two official exports use either Date/Number/Time or Day/Time.
        if len(rows) < 2:
            warnings.append(f"Лист «{sheet}»: недостаточно строк для расписания.")
            continue
        if rows and rows[0][1:3] == ["Номер", "Время"]:
            rows[0][0] = "Дата"
        elif rows:
            header = next(
                (r for r in range(min(5, len(rows))) if rows[r][:2] == ["День недели", "Время"]),
                None,
            )
            if header is not None:
                rows = [row[:1] + [""] + row[1:] for row in rows[header:]]
                merges = [
                    (r1 - header, r2 - header, c1 + (c1 >= 1), c2 + (c2 > 1))
                    for r1, r2, c1, c2 in merges
                    if r1 >= header
                ]
                rows[0][:3] = ["Дата", "Номер", "Время"]
                for c in range(3, len(rows[1]), 2):
                    rows[1][c] = "1"
        if not rows or rows[0][:3] != ["Дата", "Номер", "Время"]:
            warnings.append(
                f"Лист «{sheet}»: ручной формат, нужна отдельная проверка; автоматический импорт недоступен."
            )
            continue
        # Carry merged values only within their actual merge ranges, never arbitrary blanks.
        filled = [row[:] for row in rows]
        for r1, r2, c1, c2 in merges:
            for r in range(r1, min(r2, len(rows))):
                for c in range(c1, min(c2, len(rows[r]))):
                    filled[r][c] = rows[r1][c1]
        starts = [c for c in range(3, len(rows[0])) if rows[0][c]]
        for idx, c in enumerate(starts):
            end = starts[idx + 1] if idx + 1 < len(starts) else len(rows[0])
            name = rows[0][c]
            subs = []
            for col in range(c, end, 2):
                label = filled[1][col]
                try:
                    sg = int(float(label))
                except (ValueError, TypeError):
                    sg = 1
                subs.append((col, sg))
            result = []
            day = None
            slotrow = None
            start = endtime = None
            unparsed = []
            for r in range(2, len(rows)):
                dayname = filled[r][0].strip().lower()
                if dayname in DAYS:
                    day = DAYS[dayname]
                if rows[r][2]:
                    matches = re.findall(r"(?<!\d)(\d{1,2}):(\d{2})(?::\d{2})?", rows[r][2])
                    if len(matches) != 2:
                        start = None
                        continue
                    start, endtime = ["%02d:%s" % (int(h), m) for h, m in matches]
                    slotrow = r
                if day is None or start is None or slotrow is None:
                    continue
                offset = r - slotrow
                if offset not in (0, 1):
                    continue
                week = "odd" if offset == 0 else "even"
                for col, sg in subs:
                    raw = filled[r][col].strip().replace("\\n", "\n")
                    if not raw:
                        continue
                    room = filled[r][col + 1].strip() if col + 1 < len(filled[r]) else ""
                    override = re.match(
                        r"^с\s+(\d{1,2}:\d{2})\s+до\s+(\d{1,2}:\d{2})\s+", raw, re.I
                    )
                    lesson_start, lesson_end = (
                        (override.group(1).zfill(5), override.group(2).zfill(5))
                        if override
                        else (start, endtime)
                    )
                    cleaned = raw[override.end() :] if override else raw
                    parts = cleaned.splitlines()
                    first = parts[0]
                    kind = re.search(
                        r"\((Лекционные|Практические|Лабораторные|Семинарские)\)", first, re.I
                    )
                    if not kind:
                        unparsed.append(
                            {
                                "row": r + 1,
                                "column": col + 1,
                                "text": raw,
                                "reason": "Не распознан тип занятия",
                            }
                        )
                        continue
                    subject = first[: kind.start()].strip()
                    subject = re.sub(r"^\d+(?:\s*,\s*\d+)*\s*п\.г\.\s*", "", subject)
                    teacher = " ".join(parts[1:]).strip()
                    typ = {
                        "лекционные": "Лекция",
                        "практические": "Практика",
                        "лабораторные": "Лабораторная",
                        "семинарские": "Семинар",
                    }[kind.group(1).lower()]
                    if re.search(r"\d{1,2}[./]\d{1,2}|\b(?:с|до)\s*\d|\d+\s*нед", cleaned, re.I):
                        unparsed.append(
                            {
                                "row": r + 1,
                                "column": col + 1,
                                "text": raw,
                                "reason": "Есть ограничения по датам/неделям; нужна ручная проверка",
                            }
                        )
                        continue
                    result.append(
                        {
                            "day": day,
                            "start": lesson_start,
                            "end": lesson_end,
                            "subject": subject,
                            "type": typ,
                            "teacher": teacher[:160],
                            "room": room[:100],
                            "week": week,
                            "subgroups": [sg],
                            "note": raw,
                            "sourceCell": f"{sheet}!R{r + 1}C{col + 1}",
                        }
                    )
            # Merge identical subgroup records, then identical odd/even records.
            merged = {}
            for lesson in result:
                key = tuple(
                    lesson[k]
                    for k in ["day", "start", "end", "subject", "type", "teacher", "room", "week"]
                )
                if key in merged:
                    merged[key]["subgroups"] = sorted(
                        set(merged[key]["subgroups"] + lesson["subgroups"])
                    )
                else:
                    merged[key] = lesson
            compact = {}
            for lesson in merged.values():
                key = tuple(
                    lesson[k] for k in ["day", "start", "end", "subject", "type", "teacher", "room"]
                ) + (tuple(lesson["subgroups"]),)
                if key in compact:
                    compact[key]["week"] = "all"
                else:
                    compact[key] = lesson
            lessons = list(compact.values())
            for lesson in lessons:
                lesson["id"] = (
                    "misis-"
                    + hashlib.sha256(
                        json.dumps(
                            [
                                name,
                                lesson["day"],
                                lesson["start"],
                                lesson["end"],
                                lesson["subject"],
                                lesson["week"],
                                lesson["subgroups"],
                            ],
                            ensure_ascii=False,
                        ).encode()
                    ).hexdigest()[:20]
                )
            groups.append(
                {
                    "name": name,
                    "sheet": sheet,
                    "subgroups": sorted(set(sg for _, sg in subs)),
                    "lessons": lessons,
                    "unparsed": unparsed,
                }
            )
    return {"groups": groups, "warnings": warnings}


def sync(page_file=None, source=None, cache_dir=None):
    """Collect independent institute files; one failure must not discard other results."""
    parser = Links()
    parser.feed(Path(page_file).read_text() if page_file else download(PAGE).decode("utf-8"))
    links = parser.links
    if not links:
        raise ValueError("На странице МИСИС не найдены файлы расписания")

    def process(item):
        key = hashlib.sha256(item["url"].encode()).hexdigest()[:16]
        out = {**item, "id": key}
        try:
            cache = Path(cache_dir) / item["url"].rsplit("/", 1)[-1] if cache_dir else None
            data = cache.read_bytes() if cache and cache.exists() else download(item["url"])
            if cache:
                cache.parent.mkdir(parents=True, exist_ok=True)
                cache.write_bytes(data)
            out.update(parse_workbook(data))
            out["fetchedAt"] = datetime.datetime.now(datetime.timezone.utc).isoformat()
        except Exception as exc:
            out["groups"] = []
            out["error"] = f"Не удалось прочитать файл: {type(exc).__name__}"
        return out

    if source:
        links = [x for x in links if x["url"] == source]
    with concurrent.futures.ThreadPoolExecutor(max_workers=4) as pool:
        sources = list(pool.map(process, links))
    return {
        "page": PAGE,
        "fetchedAt": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "sources": sources,
        "weekNote": "Верхняя строка пары принята за числитель, нижняя — за знаменатель. Проверьте соответствие первой неделе семестра.",
    }


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--file")
    p.add_argument("--page-file")
    p.add_argument("--source")
    p.add_argument("--cache-dir")
    p.add_argument("--output")
    args = p.parse_args()
    try:
        data = (
            parse_workbook(Path(args.file).read_bytes())
            if args.file
            else sync(args.page_file, args.source, args.cache_dir)
        )
        result = json.dumps(data, ensure_ascii=False)
        if args.output:
            Path(args.output).write_text(result, encoding="utf-8")
        else:
            print(result)
    except Exception as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False))
        sys.exit(1)
